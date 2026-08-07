"""
Shared report utilities — Mind Mood AI Test Suites
"""
import datetime, os, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ILLEGAL = re.compile(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]")

def safe(v):
    return ILLEGAL.sub("", str(v)) if v is not None else ""

COLORS = {
    "hdr_bg":   "1E3A5F",
    "title_bg": "0F2544",
    "pass_bg":  "D1FAE5",
    "pass_fg":  "065F46",
    "fail_bg":  "FEE2E2",
    "fail_fg":  "991B1B",
    "zebra":    "F0F4F8",
    "kpi_bg":   "EEF2F6",
    "white":    "FFFFFF",
    "text":     "1E293B",
}

def _side(): return Side(style="thin", color="CBD5E1")
def _border(): return Border(left=_side(), right=_side(), top=_side(), bottom=_side())
def _fill(hex_color): return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
def _font(bold=False, color="1E293B", size=10): return Font(name="Calibri", size=size, bold=bold, color=color)
def _align(h="center"): return Alignment(horizontal=h, vertical="center", wrap_text=True)

def generate_excel_report(results, suite_name, description, out_path):
    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    total  = len(results)
    passed = sum(1 for r in results if r.get("status") == "Pass")
    failed = total - passed
    rate   = f"{passed/total*100:.1f}%" if total else "0%"

    wb = openpyxl.Workbook()

    # ── Sheet 1 : Summary ───────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary Dashboard"
    ws.sheet_view.showGridLines = False

    # Title banner
    ws.merge_cells("A1:J2")
    ws["A1"].value      = f"  Mind Mood AI  ▸  {suite_name}"
    ws["A1"].font       = Font(name="Calibri", size=16, bold=True, color=COLORS["white"])
    ws["A1"].fill       = _fill(COLORS["title_bg"])
    ws["A1"].alignment  = _align("left")
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22

    # Meta
    for r, label, value in [
        (4, "Description:", description),
        (5, "Generated:", ts),
        (6, "Environment:", "GitHub Actions CI  /  Vercel Production"),
    ]:
        ws.cell(r, 1).value = label; ws.cell(r, 1).font = _font(bold=True)
        ws.cell(r, 2).value = value; ws.cell(r, 2).font = _font()
        ws.merge_cells(f"B{r}:J{r}")

    # KPI boxes
    kpis = [("TOTAL TESTS", str(total), "A"), ("PASSED", str(passed), "D"), ("PASS RATE", rate, "G")]
    for label, val, col in kpis:
        ws.merge_cells(f"{col}8:{chr(ord(col)+1)}8")
        ws.merge_cells(f"{col}9:{chr(ord(col)+1)}10")
        lc = ws[f"{col}8"]; lc.value = label
        lc.font = _font(bold=True, color=COLORS["white"]); lc.fill = _fill(COLORS["hdr_bg"])
        lc.alignment = _align()
        vc = ws[f"{col}9"]; vc.value = val
        vc.font = Font(name="Calibri", size=20, bold=True, color=COLORS["hdr_bg"])
        vc.fill = _fill(COLORS["kpi_bg"]); vc.alignment = _align()

    # Module breakdown header
    ws.merge_cells("A12:J12")
    ws["A12"].value = "  Module / Category Breakdown"
    ws["A12"].font  = _font(bold=True, size=11, color=COLORS["white"])
    ws["A12"].fill  = _fill(COLORS["hdr_bg"])
    ws["A12"].alignment = _align("left")

    hcols = ["Module", "Tests", "Passed", "Failed", "Pass Rate"]
    hwidths = [45, 8, 8, 8, 12]
    for ci, (h, w) in enumerate(zip(hcols, hwidths), 1):
        c = ws.cell(13, ci)
        c.value = h; c.font = _font(bold=True, color=COLORS["white"])
        c.fill = _fill(COLORS["title_bg"]); c.border = _border(); c.alignment = _align()
        ws.column_dimensions[get_column_letter(ci)].width = w

    modules = {}
    for r in results:
        modules.setdefault(r.get("module","General"), []).append(r)

    for ri, (mod, cases) in enumerate(modules.items(), 14):
        p = sum(1 for r in cases if r.get("status") == "Pass")
        row_data = [mod, len(cases), p, len(cases)-p, f"{p/len(cases)*100:.0f}%"]
        for ci, val in enumerate(row_data, 1):
            c = ws.cell(ri, ci); c.value = val; c.border = _border()
            c.alignment = _align("left" if ci == 1 else "center")
            if ri % 2 == 0: c.fill = _fill(COLORS["zebra"])
            if ci == 5:
                c.font = _font(bold=True, color=COLORS["pass_fg"])
                c.fill = _fill(COLORS["pass_bg"])

    # ── Sheet 2 : Full Detail ───────────────────────────────────────────
    wd = wb.create_sheet("Test Execution Details")
    wd.sheet_view.showGridLines = False
    cols = ["Test ID","Module","Test Case Name","Preconditions","Test Steps",
            "Expected Result","Actual Result","Status","Exec Time (ms)","Timestamp"]
    widths = [12, 28, 42, 32, 55, 38, 38, 9, 14, 20]

    for ci, (h, w) in enumerate(zip(cols, widths), 1):
        c = wd.cell(1, ci)
        c.value = h; c.font = _font(bold=True, color=COLORS["white"])
        c.fill = _fill(COLORS["hdr_bg"]); c.border = _border(); c.alignment = _align()
        wd.column_dimensions[get_column_letter(ci)].width = w
    wd.row_dimensions[1].height = 30

    for ri, r in enumerate(results, 2):
        vals = [
            safe(r.get("id")),          safe(r.get("module")),
            safe(r.get("name")),        safe(r.get("preconditions")),
            safe("\n".join(r.get("steps", []))),
            safe(r.get("expected")),    safe(r.get("actual")),
            safe(r.get("status","Pass")), safe(r.get("execution_time","<1ms")),
            safe(r.get("timestamp", ts)),
        ]
        for ci, val in enumerate(vals, 1):
            c = wd.cell(ri, ci); c.value = val; c.border = _border()
            c.alignment = _align() if ci in (1,8,9,10) else _align("left")
            c.font = _font()
            if ri % 2 == 0 and ci != 8: c.fill = _fill(COLORS["zebra"])
        sc = wd.cell(ri, 8)
        if sc.value == "Pass":
            sc.font = _font(bold=True, color=COLORS["pass_fg"])
            sc.fill = _fill(COLORS["pass_bg"])
        else:
            sc.font = _font(bold=True, color=COLORS["fail_fg"])
            sc.fill = _fill(COLORS["fail_bg"])

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    try:
        wb.save(out_path)
        print(f"[Report] Saved -> {os.path.abspath(out_path)}")
    except Exception as e:
        alt = out_path.replace(".xlsx", "_alt.xlsx")
        wb.save(alt)
        print(f"[Report] Saved (alt) -> {os.path.abspath(alt)}")
